import streamlit as st
import pandas as pd
import numpy as np
import re
import plotly.express as px
import plotly.graph_objects as go
from openpyxl import load_workbook
import io

# ============================================================
# PAGE CONFIG
# ============================================================
st.set_page_config(
    page_title="Defect Management Dashboard - Conveyors",
    layout="wide",
    page_icon="📊"
)

st.markdown("""
<style>
    .stMetric { background-color: #f8f9fa; border-radius: 10px; padding: 10px; border: 1px solid #e0e0e0; }
    h1, h2, h3 { color: #1f2c4c; }
    table { font-size: 14px; }
</style>
""", unsafe_allow_html=True)

st.title("📊 Defect Management Dashboard - Conveyors")
st.caption("Upload your Excel file to view the executive summary and the detailed analysis.")

# ============================================================
# FILE UPLOAD
# ============================================================
st.sidebar.header("📁 Upload File")
uploaded_file = st.sidebar.file_uploader("Select the Excel file (.xlsx)", type=["xlsx", "xls"])

if uploaded_file is None:
    st.info("👆 Upload an Excel file from the sidebar to get started.")
    st.stop()

file_bytes = uploaded_file.read()

@st.cache_data
def get_sheet_names(_bytes):
    return pd.ExcelFile(io.BytesIO(_bytes)).sheet_names

@st.cache_data
def load_data(_bytes, sheet_name):
    df = pd.read_excel(io.BytesIO(_bytes), sheet_name=sheet_name)
    df.columns = [str(c).strip() for c in df.columns]
    for col in ["Date Open", "Date Closed"]:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce")
    return df

# ============================================================
# READ EXECUTIVE SUMMARY ("Dashboard" sheet)
# ============================================================
def find_cell(ws, target):
    target = str(target).strip().lower()
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().lower() == target:
                return cell.row, cell.column
    return None

def get_value(ws, r, c):
    return ws.cell(row=r, column=c).value

@st.cache_data
def parse_dashboard_sheet(_bytes, sheet_name="Dashboard"):
    try:
        wb = load_workbook(io.BytesIO(_bytes), data_only=True)
    except Exception:
        return None
    if sheet_name not in wb.sheetnames:
        return None
    ws = wb[sheet_name]
    result = {}

    for label, key in [("OPEN DEFECTS", "open"), ("PENDING CONFIRMATION", "pending"), ("CLOSED DEFECTS", "closed")]:
        pos = find_cell(ws, label)
        if pos:
            r, c = pos
            result[key] = get_value(ws, r + 1, c)

    pos = find_cell(ws, "Status / Severity")
    if pos:
        r, c = pos
        rows = ["Open", "Confirmation Pending", "Closed", "Total"]
        data = []
        for i, row_label in enumerate(rows, start=1):
            vals = [get_value(ws, r + i, c + j) for j in range(1, 5)]
            data.append([row_label] + vals)
        result["severity_table"] = pd.DataFrame(data, columns=["Status/Severity", "Minor", "Major", "Critical", "Total"])

    pos = find_cell(ws, "Element Type")
    if pos:
        r, c = pos
        rows_data, i = [], 2
        while True:
            label = get_value(ws, r + i, c)
            if label is None or str(label).strip() == "":
                break
            vals = [get_value(ws, r + i, c + j) for j in range(1, 7)]
            rows_data.append([label] + vals)
            i += 1
        result["element_table"] = pd.DataFrame(
            rows_data, columns=["Element Type", "Open", "Closed", "Pending", "Minor", "Major", "Critical"]
        )

    pos = find_cell(ws, "Week")
    if pos:
        r, c = pos
        weeks, i = [], 1
        while True:
            wk = get_value(ws, r + i, c)
            if wk is None or str(wk).strip() == "":
                break
            weeks.append([wk, get_value(ws, r + i, c + 1), get_value(ws, r + i, c + 2), get_value(ws, r + i, c + 3)])
            i += 1
        result["week_table"] = pd.DataFrame(weeks, columns=["Week", "Closed", "Confirmation Pending", "Open"])

    return result if result else None

def compute_fallback_summary(df):
    df_tmp = df.copy()
    df_tmp["StatusNorm"] = df_tmp["Status"].astype(str).apply(
        lambda x: "Open" if x.strip().lower() == "open"
        else ("Closed" if x.strip().lower() == "closed" else "Confirmation Pending")
    )
    open_v = (df_tmp["StatusNorm"] == "Open").sum()
    pending_v = (df_tmp["StatusNorm"] == "Confirmation Pending").sum()
    closed_v = (df_tmp["StatusNorm"] == "Closed").sum()

    sev_order, status_order = ["Minor", "Major", "Critical"], ["Open", "Confirmation Pending", "Closed"]
    pivot = pd.crosstab(df_tmp["StatusNorm"], df_tmp["Severity"]).reindex(
        index=status_order, columns=sev_order, fill_value=0
    )
    pivot["Total"] = pivot.sum(axis=1)
    total_row = pivot.sum(axis=0)
    total_row.name = "Total"
    pivot = pd.concat([pivot, pd.DataFrame([total_row])])
    df_sev = pivot.reset_index().rename(columns={"index": "Status/Severity"})

    elem_rows = []
    for etype, g in df_tmp.groupby("Type"):
        elem_rows.append([
            etype, (g["StatusNorm"] == "Open").sum(), (g["StatusNorm"] == "Closed").sum(),
            (g["StatusNorm"] == "Confirmation Pending").sum(),
            (g["Severity"] == "Minor").sum(), (g["Severity"] == "Major").sum(), (g["Severity"] == "Critical").sum()
        ])
    df_elem = pd.DataFrame(elem_rows, columns=["Element Type", "Open", "Closed", "Pending", "Minor", "Major", "Critical"])
    total = ["TOTAL"] + [df_elem[c].sum() for c in ["Open", "Closed", "Pending", "Minor", "Major", "Critical"]]
    df_elem = pd.concat([df_elem, pd.DataFrame([total], columns=df_elem.columns)], ignore_index=True)

    return {"open": open_v, "pending": pending_v, "closed": closed_v,
            "severity_table": df_sev, "element_table": df_elem, "week_table": None}

# ============================================================
# HELPERS
# ============================================================
def fmt_num(v):
    """Format numeric cell: blank if NaN/None, integer without decimals if whole number."""
    if v is None:
        return ""
    if isinstance(v, float):
        if pd.isna(v):
            return ""
        if v.is_integer():
            return str(int(v))
        return f"{v:.1f}"
    return str(v)

def parse_week_number(w):
    m = re.search(r"(\\d+)", str(w))
    return int(m.group(1)) if m else None

# ============================================================
# EXECUTIVE SUMMARY VISUAL COMPONENTS
# ============================================================
def kpi_card(title, value, color):
    st.markdown(f"""
    <div style="border:1px solid #d0d5dc; border-radius:6px; padding:15px; text-align:center; background-color:#f8f9fb;">
        <div style="font-size:13px; color:#5a6474; font-weight:600; letter-spacing:1px;">{title}</div>
        <div style="font-size:36px; font-weight:bold; color:{color};">{value}</div>
    </div>
    """, unsafe_allow_html=True)

STATUS_TEXT_COLOR = {
    "Open": "#b03a2e",
    "Confirmation Pending": "#9a6a06",
    "Closed": "#1e7e45",
    "Total": "#1f2c4c"
}

def render_severity_table(df_sev):
    html = "<table style='width:100%; border-collapse:collapse; text-align:center; border:1px solid #c9ced6;'>"
    html += "<tr>" + "".join(
        f"<th style='background-color:#1f2c4c;color:#ffffff;padding:8px;border:1px solid #c9ced6;font-weight:600;'>{c}</th>"
        for c in df_sev.columns) + "</tr>"
    for idx, row in df_sev.iterrows():
        bg = "#f4f6f9" if idx % 2 == 0 else "#ffffff"
        html += f"<tr style='background-color:{bg};'>"
        for i, val in enumerate(row):
            col = df_sev.columns[i]
            if col == "Status/Severity":
                color = STATUS_TEXT_COLOR.get(str(val).strip(), "#2c3e50")
                weight = "700"
                display_val = val
            else:
                color = "#2c3e50"
                weight = "700" if str(row["Status/Severity"]).strip() == "Total" else "500"
                display_val = fmt_num(val)
            html += f"<td style='padding:8px;border:1px solid #c9ced6;color:{color};font-weight:{weight};'>{display_val}</td>"
        html += "</tr>"
    html += "</table>"
    st.markdown(html, unsafe_allow_html=True)

def render_element_table(df_elem):
    html = "<table style='width:100%; border-collapse:collapse; text-align:center; border:1px solid #c9ced6;'>"
    html += ("<tr><th style='background-color:#1f2c4c;color:#ffffff;padding:8px;border:1px solid #c9ced6;font-weight:600;'>Element Type</th>"
              "<th colspan='3' style='background-color:#2f3f5c;color:#ffffff;padding:8px;border:1px solid #c9ced6;font-weight:600;'>STATUS BREAKDOWN</th>"
              "<th colspan='3' style='background-color:#2f3f5c;color:#ffffff;padding:8px;border:1px solid #c9ced6;font-weight:600;'>SEVERITY BREAKDOWN</th></tr>")
    html += "<tr><th style='padding:8px;border:1px solid #c9ced6;background-color:#e4e8ee;'></th>"
    for c in ["Open", "Closed", "Pending", "Minor", "Major", "Critical"]:
        html += f"<th style='padding:8px;border:1px solid #c9ced6;background-color:#e4e8ee;color:#1f2c4c;font-weight:600;'>{c}</th>"
    html += "</tr>"
    for idx, row in df_elem.iterrows():
        is_total = str(row["Element Type"]).strip().upper() == "TOTAL"
        bg = "#eef1f6" if is_total else ("#f4f6f9" if idx % 2 == 0 else "#ffffff")
        w = "700" if is_total else "500"
        html += f"<tr style='background-color:{bg};'>"
        html += f"<td style='padding:8px;border:1px solid #c9ced6;font-weight:{w};color:#1f2c4c;text-align:left;padding-left:12px;'>{row['Element Type']}</td>"
        for c in ["Open", "Closed", "Pending", "Minor", "Major", "Critical"]:
            html += f"<td style='padding:8px;border:1px solid #c9ced6;font-weight:{w};color:#2c3e50;'>{fmt_num(row[c])}</td>"
        html += "</tr>"
    html += "</table>"
    st.markdown(html, unsafe_allow_html=True)

def render_burndown_chart(df_week, key="burndown_default"):
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=df_week["Week"], y=df_week["Closed"], name="Closed", mode="lines",
                              stackgroup="one", fillcolor="rgba(46,139,87,0.75)", line=dict(color="#2e8b57")))
    fig.add_trace(go.Scatter(x=df_week["Week"], y=df_week["Confirmation Pending"], name="Confirmation Pending",
                              mode="lines", stackgroup="one", fillcolor="rgba(217,164,6,0.75)", line=dict(color="#d9a406")))
    fig.add_trace(go.Scatter(x=df_week["Week"], y=df_week["Open"], name="Open", mode="lines",
                              stackgroup="one", fillcolor="rgba(176,58,46,0.75)", line=dict(color="#b03a2e")))
    fig.update_layout(
        title="Defect Status Trend & Burndown per Calendar Week",
        xaxis_title="Week", yaxis_title="Count",
        legend=dict(orientation="h", y=-0.2), height=430
    )
    st.plotly_chart(fig, use_container_width=True, key=key)

def render_category_table(df_cat):
    html = "<table style='width:100%; border-collapse:collapse; text-align:center; border:1px solid #c9ced6;'>"
    html += "<tr>" + "".join(
        f"<th style='background-color:#1f2c4c;color:#ffffff;padding:8px;border:1px solid #c9ced6;font-weight:600;'>{c}</th>"
        for c in df_cat.columns) + "</tr>"
    for idx, row in df_cat.iterrows():
        is_total = str(row["Category"]).strip().upper() == "TOTAL"
        bg = "#eef1f6" if is_total else ("#f4f6f9" if idx % 2 == 0 else "#ffffff")
        w = "700" if is_total else "500"
        html += f"<tr style='background-color:{bg};'>"
        html += f"<td style='padding:8px;border:1px solid #c9ced6;font-weight:{w};color:#1f2c4c;text-align:left;padding-left:12px;'>{row['Category']}</td>"
        for c in df_cat.columns[1:]:
            val = row[c]
            display_val = f"{val:.1f}%" if c == "% of Total" else fmt_num(val)
            color = "#b03a2e" if c == "Critical" and isinstance(val, (int, float)) and val > 0 else "#2c3e50"
            html += f"<td style='padding:8px;border:1px solid #c9ced6;font-weight:{w};color:{color};'>{display_val}</td>"
        html += "</tr>"
    html += "</table>"
    st.markdown(html, unsafe_allow_html=True)

# ============================================================
# FORECAST LOGIC
# ============================================================
def compute_forecast(df_week, method="auto", manual_rate=None):
    """
    Projects the burndown forward assuming no new defects are raised.
    Returns: (df_actual_with_weeknum, df_forecast_or_None, weeks_needed_or_None, rate)
    """
    df_w = df_week.copy()
    df_w["WeekNum"] = df_w["Week"].apply(parse_week_number)
    df_w = df_w.dropna(subset=["WeekNum"]).sort_values("WeekNum").reset_index(drop=True)
    if df_w.empty or len(df_w) < 2:
        return df_w, None, None, None

    last_row = df_w.iloc[-1]
    backlog = last_row["Open"] + last_row["Confirmation Pending"]
    total = backlog + last_row["Closed"]

    if backlog <= 0:
        return df_w, None, 0, None

    if method == "manual" and manual_rate is not None:
        rate = manual_rate
    else:
        x = df_w["WeekNum"].values.astype(float)
        y = df_w["Closed"].values.astype(float)
        slope, _ = np.polyfit(x, y, 1)
        rate = slope

    if rate is None or rate <= 0:
        return df_w, None, None, rate

    weeks_needed = int(np.ceil(backlog / rate))
    open_ratio = last_row["Open"] / backlog if backlog > 0 else 0
    pending_ratio = last_row["Confirmation Pending"] / backlog if backlog > 0 else 0
    last_week_num = int(last_row["WeekNum"])

    forecast_rows = []
    for i in range(1, weeks_needed + 1):
        remaining_backlog = max(0.0, backlog - rate * i)
        cum_closed = min(total, last_row["Closed"] + rate * i)
        forecast_rows.append({
            "Week": f"CW{last_week_num + i}",
            "Closed": cum_closed,
            "Confirmation Pending": remaining_backlog * pending_ratio,
            "Open": remaining_backlog * open_ratio,
        })
    df_forecast = pd.DataFrame(forecast_rows)
    return df_w, df_forecast, weeks_needed, rate

def render_burndown_with_forecast(df_actual, df_forecast, weeks_needed, key="burndown_forecast_default"):
    fig = go.Figure()
    # Actual (solid stacked areas)
    fig.add_trace(go.Scatter(x=df_actual["Week"], y=df_actual["Closed"], name="Closed (Actual)", mode="lines",
                              stackgroup="actual", fillcolor="rgba(46,139,87,0.75)", line=dict(color="#2e8b57")))
    fig.add_trace(go.Scatter(x=df_actual["Week"], y=df_actual["Confirmation Pending"], name="Confirmation Pending (Actual)",
                              mode="lines", stackgroup="actual", fillcolor="rgba(217,164,6,0.75)", line=dict(color="#d9a406")))
    fig.add_trace(go.Scatter(x=df_actual["Week"], y=df_actual["Open"], name="Open (Actual)", mode="lines",
                              stackgroup="actual", fillcolor="rgba(176,58,46,0.75)", line=dict(color="#b03a2e")))

    # Bridge point so forecast lines connect visually with the last actual week
    bridge = pd.DataFrame([{
        "Week": df_actual["Week"].iloc[-1],
        "Closed": df_actual["Closed"].iloc[-1],
        "Confirmation Pending": df_actual["Confirmation Pending"].iloc[-1],
        "Open": df_actual["Open"].iloc[-1],
    }])
    df_fc_plot = pd.concat([bridge, df_forecast], ignore_index=True)

    fig.add_trace(go.Scatter(
        x=df_fc_plot["Week"], y=df_fc_plot["Closed"],
        name="Closed (Forecast)", mode="lines", line=dict(color="#2e8b57", dash="dash")
    ))
    fig.add_trace(go.Scatter(
        x=df_fc_plot["Week"], y=df_fc_plot["Closed"] + df_fc_plot["Confirmation Pending"],
        name="Closed + Pending (Forecast)", mode="lines", line=dict(color="#d9a406", dash="dash")
    ))
    fig.add_trace(go.Scatter(
        x=df_fc_plot["Week"], y=df_fc_plot["Closed"] + df_fc_plot["Confirmation Pending"] + df_fc_plot["Open"],
        name="Total Backlog (Forecast)", mode="lines", line=dict(color="#b03a2e", dash="dash")
    ))

    fig.add_vline(x=df_actual["Week"].iloc[-1], line_width=1, line_dash="dot", line_color="#888888")

    title = "Defect Status Trend & Burndown per Calendar Week"
    if weeks_needed:
        title += f"  —  Projected full closure in ~{weeks_needed} more week(s)"

    fig.update_layout(title=title, xaxis_title="Week", yaxis_title="Count",
                       legend=dict(orientation="h", y=-0.3), height=460)
    st.plotly_chart(fig, use_container_width=True, key=key)

# ============================================================
# SECTION 1: EXECUTIVE SUMMARY (main part)
# ============================================================
st.header("📌 Executive Summary")

sheet_names = get_sheet_names(file_bytes)
summary = parse_dashboard_sheet(file_bytes, "Dashboard") if "Dashboard" in sheet_names else None

fallback_used = False
if summary is None:
    default_sheet_fb = "Data" if "Data" in sheet_names else sheet_names[0]
    df_for_fallback = load_data(file_bytes, default_sheet_fb)
    if all(c in df_for_fallback.columns for c in ["Type", "Status", "Severity"]):
        summary = compute_fallback_summary(df_for_fallback)
        fallback_used = True

if summary:
    if fallback_used:
        st.warning("The **Dashboard** sheet could not be found (or interpreted). "
                    "An equivalent summary was computed from the **Data** sheet instead. "
                    "The weekly trend chart and forecast are not available in this mode.")

    c1, c2, c3 = st.columns(3)
    with c1: kpi_card("OPEN DEFECTS", summary.get("open", "N/A"), "#b03a2e")
    with c2: kpi_card("PENDING CONFIRMATION", summary.get("pending", "N/A"), "#9a6a06")
    with c3: kpi_card("CLOSED DEFECTS", summary.get("closed", "N/A"), "#1e7e45")

    st.markdown("####")
    left, right = st.columns([1, 1.3])

    with left:
        st.subheader("Status / Severity")
        if summary.get("severity_table") is not None:
            render_severity_table(summary["severity_table"])
        st.markdown("####")
        st.subheader("Element Type")
        if summary.get("element_table") is not None:
            render_element_table(summary["element_table"])

    with right:
        if summary.get("week_table") is not None:
            render_burndown_chart(summary["week_table"], key="burndown_summary")
        else:
            st.info("The weekly trend (Burndown) chart requires the **Dashboard** sheet "
                    "with the weekly table (Week / Closed / Confirmation Pending / Open) from the original Excel file.")

    # -------------------- CLOSURE FORECAST --------------------
    st.markdown("---")
    st.subheader("🔮 Closure Forecast")
    st.caption("Projects how long it would take to close the remaining backlog, "
               "assuming **no new defects are raised** and the closure pace stays constant.")

    if summary.get("week_table") is not None:
        show_forecast = st.checkbox("Show closure forecast", value=False)

        if show_forecast:
            method_label = st.radio(
                "Forecast method",
                ["Automatic (based on historical closure trend)", "Manual (enter a weekly closure rate)"],
                horizontal=True
            )
            manual_rate = None
            if method_label.startswith("Manual"):
                manual_rate = st.number_input(
                    "Estimated defects closed per week", min_value=0.0, value=10.0, step=1.0
                )

            df_actual, df_forecast, weeks_needed, rate = compute_forecast(
                summary["week_table"],
                method="manual" if manual_rate else "auto",
                manual_rate=manual_rate
            )

            if weeks_needed == 0:
                st.success("✅ The backlog is already at zero — there is nothing left to close.")
                render_burndown_chart(summary["week_table"], key="burndown_zero_backlog")
            elif df_forecast is None:
                st.warning(
                    "⚠️ Based on the historical trend, the closure rate is zero or negative. "
                    "At the current pace, the remaining backlog would not be resolved. "
                    "Try the **Manual** option above to test a different weekly closure rate."
                )
                render_burndown_chart(summary["week_table"], key="burndown_no_forecast")
            else:
                last_actual = df_actual.iloc[-1]
                remaining_backlog = int(last_actual["Open"] + last_actual["Confirmation Pending"])
                projected_week = df_forecast["Week"].iloc[-1]

                cf1, cf2, cf3 = st.columns(3)
                with cf1: kpi_card("REMAINING BACKLOG", remaining_backlog, "#b03a2e")
                with cf2: kpi_card("ESTIMATED WEEKS TO CLOSE", weeks_needed, "#9a6a06")
                with cf3: kpi_card("PROJECTED CLOSURE WEEK", projected_week, "#1e7e45")

                render_burndown_with_forecast(df_actual, df_forecast, weeks_needed, key="burndown_forecast_active")
                st.caption(
                    f"Closure rate used: **{rate:.1f} defects/week**. "
                    "Dashed lines represent the projected scenario. The split between Open and "
                    "Confirmation Pending in the forecast keeps the same proportion observed in the last actual week."
                )
    else:
        st.info("The forecast requires the **Dashboard** sheet with the weekly table from the original Excel file.")

else:
    st.error("Unable to generate the executive summary from this file. Please verify it contains "
              "a **Dashboard** sheet or a **Data** sheet with the required columns.")

st.markdown("---")

# ============================================================
# SECTION 2: DETAILED ANALYSIS (Data sheet)
# ============================================================
st.header("🔍 Detailed Analysis")

default_sheet = "Data" if "Data" in sheet_names else sheet_names[0]
sheet_selected = st.sidebar.selectbox(
    "Sheet to analyze (Detailed Analysis)", sheet_names, index=sheet_names.index(default_sheet)
)
df_raw = load_data(file_bytes, sheet_selected)

required_cols = ["Conveyor", "Station", "Type", "Category", "Status", "Severity"]
missing = [c for c in required_cols if c not in df_raw.columns]
if missing:
    st.error(f"The selected sheet is missing required columns: {missing}")
    st.stop()

df = df_raw.copy()
for col in ["Status", "Severity", "Type", "Category", "Responsible"]:
    if col in df.columns:
        df[col] = df[col].astype(str).str.strip().replace({"nan": None, "None": None})

# ---- Filters ----
st.sidebar.header("🔎 Filters (Detailed Analysis)")

def multiselect_filter(label, col):
    if col in df.columns:
        options = sorted([o for o in df[col].dropna().unique().tolist()])
        return st.sidebar.multiselect(label, options, default=options)
    return None

type_sel = multiselect_filter("Type", "Type")
category_sel = multiselect_filter("Category", "Category")
status_sel = multiselect_filter("Status", "Status")
severity_sel = multiselect_filter("Severity", "Severity")
responsible_sel = multiselect_filter("Responsible", "Responsible")

mask = pd.Series(True, index=df.index)
if type_sel is not None: mask &= df["Type"].isin(type_sel)
if category_sel is not None: mask &= df["Category"].isin(category_sel)
if status_sel is not None: mask &= df["Status"].isin(status_sel)
if severity_sel is not None: mask &= df["Severity"].isin(severity_sel)
if responsible_sel is not None and "Responsible" in df.columns:
    mask &= df["Responsible"].isin(responsible_sel) | df["Responsible"].isna()

if "Date Open" in df.columns and df["Date Open"].notna().any():
    min_date, max_date = df["Date Open"].min().date(), df["Date Open"].max().date()
    date_range = st.sidebar.date_input("Opening Date Range", value=(min_date, max_date),
                                        min_value=min_date, max_value=max_date)
    if isinstance(date_range, tuple) and len(date_range) == 2:
        start, end = date_range
        mask &= (df["Date Open"].dt.date >= start) & (df["Date Open"].dt.date <= end)

df_f = df[mask].copy()
st.sidebar.markdown("---")
st.sidebar.write(f"**Records displayed:** {len(df_f)} / {len(df)}")

COLOR_STATUS = {"Open": "#b03a2e", "Closed": "#1e7e45", "Confirmation pending": "#d9a406"}
COLOR_SEV = {"Minor": "#3a7ab0", "Major": "#d9a406", "Critical": "#b03a2e"}

c1, c2 = st.columns(2)
with c1:
    st.subheader("Status Distribution")
    status_counts = df_f["Status"].value_counts().reset_index()
    status_counts.columns = ["Status", "Count"]
    fig = px.pie(status_counts, names="Status", values="Count", hole=0.45,
                 color="Status", color_discrete_map=COLOR_STATUS)
    fig.update_traces(textinfo="percent+value")
    st.plotly_chart(fig, use_container_width=True, key="status_distribution_pie")

with c2:
    st.subheader("Severity Distribution")
    if "Severity" in df_f.columns:
        sev_counts = df_f["Severity"].value_counts().reset_index()
        sev_counts.columns = ["Severity", "Count"]
        fig2 = px.bar(sev_counts, x="Severity", y="Count", color="Severity",
                      text="Count", color_discrete_map=COLOR_SEV)
        fig2.update_layout(showlegend=False)
        st.plotly_chart(fig2, use_container_width=True, key="severity_distribution_bar")

c3, c4 = st.columns(2)
with c3:
    st.subheader("Status by Element Type")
    type_status = df_f.groupby(["Type", "Status"]).size().reset_index(name="Count")
    fig3 = px.bar(type_status, x="Type", y="Count", color="Status", barmode="stack",
                  color_discrete_map=COLOR_STATUS)
    st.plotly_chart(fig3, use_container_width=True, key="status_by_type_bar")

with c4:
    st.subheader("Top 15 Stations with Most Defects")
    if "Station" in df_f.columns:
        station_counts = df_f["Station"].value_counts().head(15).reset_index()
        station_counts.columns = ["Station", "Count"]
        fig4 = px.bar(station_counts, x="Count", y="Station", orientation="h", text="Count",
                      color="Count", color_continuous_scale="Reds")
        fig4.update_layout(yaxis={"categoryorder": "total ascending"}, coloraxis_showscale=False)
        st.plotly_chart(fig4, use_container_width=True, key="top_stations_bar")

# ============================================================
# CATEGORY ANALYSIS
# ============================================================
st.markdown("---")
st.subheader("📂 Category Analysis")
st.caption("Breakdown of defects by root-cause category (e.g. Acceleration / Mechanical, Sensors, Parametrization).")

if "Category" in df_f.columns and df_f["Category"].notna().any():
    cat_base = df_f.dropna(subset=["Category"]).copy()
    cat_base["StatusNorm"] = cat_base["Status"].apply(
        lambda x: "Open" if str(x).strip().lower() == "open"
        else ("Closed" if str(x).strip().lower() == "closed" else "Confirmation Pending")
    )

    cat_summary = cat_base.groupby("Category").agg(
        Total=("Category", "size"),
        Open=("StatusNorm", lambda s: (s == "Open").sum()),
        Confirmation_Pending=("StatusNorm", lambda s: (s == "Confirmation Pending").sum()),
        Closed=("StatusNorm", lambda s: (s == "Closed").sum()),
        Critical=("Severity", lambda s: (s == "Critical").sum())
    ).reset_index().rename(columns={"Confirmation_Pending": "Confirmation Pending"})

    cat_summary["% of Total"] = (cat_summary["Total"] / cat_summary["Total"].sum() * 100).round(1)
    cat_summary = cat_summary.sort_values("Total", ascending=False).reset_index(drop=True)

    total_row = {
        "Category": "TOTAL",
        "Total": cat_summary["Total"].sum(),
        "Open": cat_summary["Open"].sum(),
        "Confirmation Pending": cat_summary["Confirmation Pending"].sum(),
        "Closed": cat_summary["Closed"].sum(),
        "Critical": cat_summary["Critical"].sum(),
        "% of Total": 100.0
    }
    cat_summary_display = pd.concat([cat_summary, pd.DataFrame([total_row])], ignore_index=True)
    cat_summary_display = cat_summary_display[
        ["Category", "Total", "% of Total", "Open", "Confirmation Pending", "Closed", "Critical"]
    ]

    col_tbl, col_charts = st.columns([1, 1.3])

    with col_tbl:
        render_category_table(cat_summary_display)

    with col_charts:
        fig_donut = px.pie(
            cat_summary, names="Category", values="Total", hole=0.45,
            title="Defect Volume Share by Category",
            color_discrete_sequence=["#2f3f5c", "#3a7ab0", "#8fa8c4", "#c9ced6"]
        )
        fig_donut.update_traces(textinfo="percent+value")
        st.plotly_chart(fig_donut, use_container_width=True, key="category_donut")

        cat_status_melt = cat_summary.melt(
            id_vars="Category", value_vars=["Open", "Confirmation Pending", "Closed"],
            var_name="Status", value_name="Count"
        )
        fig_cat_status = px.bar(
            cat_status_melt, x="Category", y="Count", color="Status", barmode="stack",
            title="Resolution Status by Category",
            color_discrete_map={"Open": "#b03a2e", "Confirmation Pending": "#d9a406", "Closed": "#1e7e45"}
        )
        st.plotly_chart(fig_cat_status, use_container_width=True, key="category_status_bar")
else:
    st.info("The Category column is not available or has no data in the current filtered selection.")

# ============================================================
# WORKLOAD BY RESPONSIBLE
# ============================================================
if "Responsible" in df_f.columns and df_f["Responsible"].notna().any():
    st.markdown("---")
    st.subheader("👷 Workload by Responsible")
    resp = df_f.dropna(subset=["Responsible"]).groupby(["Responsible", "Status"]).size().reset_index(name="Count")
    fig6 = px.bar(resp, x="Responsible", y="Count", color="Status", barmode="stack", color_discrete_map=COLOR_STATUS)
    st.plotly_chart(fig6, use_container_width=True, key="workload_bar")

st.markdown("---")
st.subheader("📋 Record Details")
search = st.text_input("Search in Conveyor / Station / Comments")
df_show = df_f.copy()
if search:
    mask_search = pd.Series(False, index=df_show.index)
    for col in ["Conveyor", "Station", "Comments"]:
        if col in df_show.columns:
            mask_search |= df_show[col].astype(str).str.contains(search, case=False, na=False)
    df_show = df_show[mask_search]
st.dataframe(df_show, use_container_width=True, height=400)

col_dl1, col_dl2 = st.columns(2)
with col_dl1:
    csv = df_show.to_csv(index=False).encode("utf-8-sig")
    st.download_button("⬇️ Download Filtered CSV", data=csv, file_name="filtered_defects.csv", mime="text/csv")
with col_dl2:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="xlsxwriter") as writer:
        df_show.to_excel(writer, index=False, sheet_name="Data")
    st.download_button("⬇️ Download Filtered Excel", data=buffer.getvalue(), file_name="filtered_defects.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
